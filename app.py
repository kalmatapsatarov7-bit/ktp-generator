import io
import time
import streamlit as st
from docx import Document
import google.generativeai as genai
from PIL import Image

# Настройка страницы Streamlit
st.set_page_config(
    page_title="Генератор поурочных планов", page_icon="📚", layout="centered"
)

# --- РЕКЛАМНЫЙ БЛОК 1: В сайдбаре ---
st.sidebar.markdown("### 📢 Партнеры и реклама")
st.sidebar.info(
    "💡 **Место для вашей рекламы!**\n\n"
    "Хотите прорекламировать свой образовательный курс, канал или услугу? "
    "Размещайте баннеры здесь.\n\n"
    "[Написать менеджеру 👉](https://t.me/your_telegram)"
)
st.sidebar.markdown("---")

# Заголовок приложения
st.title("📚 Умный генератор поурочных планов")
st.write(
    "Создавайте поурочные планы по теме или загруженному КТП. Количество уроков подстраивается под ваши материалы автоматически!"
)

# Настройка API ключа (берётся из секрета Streamlit или вводится вручную)
if "GEMINI_API_KEY" in st.secrets:
    api_key = st.secrets["GEMINI_API_KEY"]
else:
    api_key = st.sidebar.text_input("Введите Gemini API Key", type="password")

if api_key:
    genai.configure(api_key=api_key)
else:
    st.warning(
        "Пожалуйста, укажите Gemini API Key в настройках secrets или в сайдбаре."
    )

# Используем актуальный алиас latest
model_name = "gemini-flash-latest"

# Блок ввода данных
st.subheader("1. Исходные данные для генерации")
input_option = st.radio(
    "Выберите способ ввода темы:",
    ["Текстовый ввод темы / КТП", "Загрузить фото (страница КТП / учебника)"],
)

uploaded_image = None
text_prompt = ""
custom_lessons_count = 16  # Значение по умолчанию для текстового ввода

if input_option == "Текстовый ввод темы / КТП":
    text_prompt = st.text_area(
        "Введите тему, предмет и класс (например: Русский язык, 7 класс, тема: Имя существительное):",
        placeholder="Укажите предмет, класс и основные разделы темы...",
    )
    custom_lessons_count = st.slider(
        "Количество уроков для генерации:", min_value=1, max_value=35, value=16
    )
else:
    uploaded_file = st.file_uploader(
        "Загрузите фото (страница КТП / учебника)", type=["png", "jpg", "jpeg"]
    )
    if uploaded_file is not None:
        uploaded_image = Image.open(uploaded_file)
        st.image(
            uploaded_image,
            caption="Загруженное изображение КТП",
            use_container_width=True,
        )
    extra_notes = st.text_input(
        "Дополнительные пожелания к урокам (необязательно):",
        placeholder="Например, сделать упор на практические упражнения...",
    )


# Функция генерации документов через Gemini с динамическим количеством уроков и защитой от лимитов
def generate_lessons_adaptive(prompt_content, image_obj=None, target_count=16):
    generation_config = {
        "temperature": 0.7,
        "max_output_tokens": 8192,
    }

    model = genai.GenerativeModel(
        model_name=model_name, generation_config=generation_config
    )

    if image_obj:
        base_instruction = (
            "Ты — опытный методист и учитель. Твоя задача — внимательно проанализировать "
            "загруженное изображение (КТП) и составить качественные поурочные планы "
            "строго на каждый урок, указанный в этом материале (сохрани их точное количество и последовательность). "
            "Для каждого урока распиши: Номер и тему урока, цель, основные этапы урока, "
            "краткое содержание материала и домашнее задание. Пиши структурировано."
        )
    else:
        base_instruction = (
            f"Ты — опытный методист и учитель. Твоя задача — составить ровно {target_count} подробных "
            "поурочных планов по стандартам образования. "
            "Для каждого урока распиши: Номер и тему урока, цель, основные этапы урока, "
            "краткое содержание материала и домашнее задание."
        )

    if image_obj:
        contents = [
            base_instruction,
            image_obj,
            "Извлеки точную структуру и количество уроков с этого изображения и распиши каждый урок.",
        ]
        if prompt_content:
            contents.append(f"Дополнительные пожелания: {prompt_content}")
    else:
        contents = f"{base_instruction}\n\nЗапрос/Тема: {prompt_content}"

    # Надежная защита с ожиданием при превышении лимитов бесплатного тарифа (ошибка 429)
    max_retries = 4
    for attempt in range(max_retries):
        try:
            response = model.generate_content(contents)
            return response.text
        except Exception as e:
            error_str = str(e)
            if "429" in error_str or "ResourceExhausted" in error_str:
                if attempt < max_retries - 1:
                    wait_time = 15 * (attempt + 1)
                    time.sleep(wait_time)
                    continue
            raise e
    return None


# Функция создания Word-файла в памяти
def create_docx(text_content):
    doc = Document()
    doc.add_heading("Комплект поурочных планов", 0)

    for line in text_content.split("\n"):
        if line.strip().startswith("Урок") or line.strip().startswith("#"):
            doc.add_heading(line.strip("# "), level=2)
        elif line.strip():
            doc.add_paragraph(line)

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream


# Кнопка запуска генерации
st.markdown("---")
if st.button("🚀 Сгенерировать поурочные планы", type="primary"):
    if not api_key:
        st.error("Пожалуйста, укажите API-ключ!")
    elif input_option == "Текстовый ввод темы / КТП" and not text_prompt.strip():
        st.warning("Пожалуйста, введите тему урока или КТП.")
    elif (
        input_option == "Загрузить фото (страница КТП / учебника)"
        and uploaded_image is None
    ):
        st.warning("Пожалуйста, загрузите изображение.")
    else:
        spinner_text = (
            "⏳ Анализируем КТП с фото и генерируем планы под вашу программу..."
            if uploaded_image
            else "⏳ Идет генерация подробных уроков..."
        )
        with st.spinner(spinner_text):
            try:
                result_text = generate_lessons_adaptive(
                    text_prompt if input_option == "Текстовый ввод темы / КТП" else extra_notes,
                    uploaded_image,
                    custom_lessons_count,
                )

                if result_text:
                    st.success("✅ Поурочные планы успешно готовы!")
                    st.session_state["generated_lessons"] = result_text
                else:
                    st.error(
                        "Не удалось получить ответ от модели. Попробуйте еще раз."
                    )
            except Exception as e:
                st.error(
                    f"Сервер временно перегружен или превышен лимит запросов (ошибка API). Подождите минуту и попробуйте снова. Детали: {str(e)}"
                )

# Вывод результатов и кнопки скачивания, если они есть в памяти
if "generated_lessons" in st.session_state:
    st.subheader("📖 Результат генерации:")
    st.markdown(st.session_state["generated_lessons"])

    docx_file = create_docx(st.session_state["generated_lessons"])

    st.download_button(
        label="📥 Скачать поурочные планы в формате Word (.docx)",
        data=docx_file,
        file_name="Lessons_Plan.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

# --- РЕКЛАМНЫЙ БЛОК 2: В подвале сайта (в самом низу) ---
st.markdown("---")
st.markdown(
    """
    <div style="background-color: #f0f2f6; padding: 15px; border-radius: 10px; text-align: center;">
        <p style="margin: 0; font-weight: bold; color: #31333F;">🌟 Специальное предложение / Реклама</p>
        <p style="margin: 5px 0 0 0; font-size: 14px; color: #555;">
            Полезные инструменты, методички и материалы для учителей. Подписывайтесь на наши обновления! 
            <a href="https://t.me/your_telegram" target="_blank">Узнать подробнее</a>
        </p>
    </div>
    """,
    unsafe_allow_html=True,
)
